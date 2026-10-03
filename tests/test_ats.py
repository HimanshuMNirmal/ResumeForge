import pytest
from docx import Document
from docx.shared import Inches

from helpers import load_jd, load_profile
from resume_builder.ats_check import check_file
from resume_builder.matcher import select
from resume_builder.render import docx_to_pdf, render_docx
from resume_builder.render.pdf import find_soffice
from resume_builder.textutil import SkillIndex
from resume_builder.writer import base_resume


def _good(tmp_path):
    p, jd = load_profile(), load_jd()
    r = base_resume(p, jd, select(p, jd))
    return p, jd, r, render_docx(r, tmp_path / "good.docx")


def test_good_resume_scores_high_without_pdf(tmp_path):
    p, jd, r, path = _good(tmp_path)
    rep = check_file(path, resume=r, jd=jd, idx=SkillIndex(p))
    assert rep.score("Format compliance") >= 13
    assert rep.score("Parse fidelity") >= 33
    assert "Python" in rep.matched_must and "Kubernetes" in rep.missing_must


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_good_resume_full_marks_with_pdf(tmp_path):
    p, jd, r, path = _good(tmp_path)
    pdf = docx_to_pdf(path, tmp_path / "good.pdf")
    rep = check_file(path, resume=r, jd=jd, idx=SkillIndex(p), pdf_path=pdf)
    assert rep.score("Parse fidelity") == 40
    assert rep.score("Format compliance") == 15
    assert rep.pages == 1


def test_broken_resume_is_caught(tmp_path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Jane Roe | jane@example.com"  # contact info hidden in header
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Experience"
    t.cell(0, 1).text = "Acme  2020 - 2022"
    t.cell(1, 0).text = "Skills"
    t.cell(1, 1).text = "Python"
    doc.add_paragraph("Responsible for various things")
    path = tmp_path / "bad.docx"
    doc.save(path)
    rep = check_file(path, jd=load_jd(), max_pages=1)
    names = {c.name: c for c in rep.components["Format compliance"]}
    assert names["no tables"].earned == 0
    assert names["no header/footer content"].earned == 0
    assert rep.score("Format compliance") < 10
    assert rep.score("Parse fidelity") < 25
    assert rep.overall < 60


def test_garbled_characters_detected(tmp_path):
    doc = Document()
    doc.add_paragraph("Jane Roe")
    doc.add_paragraph("jane@example.com | +91 98000 00000")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Engineer \ufb01xed \ue000 thing  Jan 2020 – Present")
    path = tmp_path / "g.docx"
    doc.save(path)
    rep = check_file(path)
    c = next(c for c in rep.components["Parse fidelity"] if c.name.startswith("no garbled"))
    assert c.earned == 0 and "U+FB01" in c.detail


def test_two_column_detected(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph("Jane Roe")
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    doc.sections[0]._sectPr.append(cols)
    path = tmp_path / "c.docx"
    doc.save(path)
    rep = check_file(path)
    c = next(c for c in rep.components["Format compliance"] if c.name == "single column")
    assert c.earned == 0


def test_keyword_stuffing_and_weak_bullets_flagged(tmp_path):
    from resume_builder.models import JD

    doc = Document()
    doc.add_paragraph("Jane Roe")
    doc.add_paragraph("jane@example.com")
    doc.add_paragraph("Experience")
    for _ in range(3):
        doc.add_paragraph("Python Python Python Python Python Python developer", style="List Bullet")
    path = tmp_path / "s.docx"
    doc.save(path)
    rep = check_file(path, jd=JD(must_have=["Python"]))
    assert any("stuffing" in w for w in rep.warnings)
    assert any("action verb" in w for w in rep.warnings)


def test_overall_ignores_keywords_without_jd(tmp_path):
    p, jd, r, path = _good(tmp_path)
    rep = check_file(path, resume=r)
    assert "Keyword coverage" in rep.not_applicable
    assert 0 < rep.overall <= 100


def test_readability_flags_unquantified_repeated_and_inconsistent_bullets(tmp_path):
    doc = Document()
    doc.add_paragraph("Jane Roe")
    doc.add_paragraph("jane@example.com")
    doc.add_paragraph("Experience")
    for t in ["Built a thing", "Built another thing.", "Built a third thing", "Designed a fourth thing"]:
        doc.add_paragraph(t, style="List Bullet")
    path = tmp_path / "r.docx"
    doc.save(path)
    rep = check_file(path)
    by = {c.name: c for c in rep.components["Readability"]}
    assert by["bullets with a number/metric (aim for >= 50%)"].earned == 0
    assert by["varied opening verbs (max 2 per verb)"].earned < 1
    assert by["consistent bullet punctuation"].earned == 0
    assert any("contain a number" in w for w in rep.warnings)


def test_rendered_bullets_have_no_trailing_periods(tmp_path):
    p, jd = load_profile(), load_jd()
    r = base_resume(p, jd, select(p, jd))
    r.experience[0].bullets[0].text += "."
    path = render_docx(r, tmp_path / "p.docx")
    from resume_builder.ats_check import extract_docx_lines
    assert not any(l.startswith("- ") and l.endswith(".") for l in extract_docx_lines(path))
