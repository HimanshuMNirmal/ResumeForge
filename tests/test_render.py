import re
import zipfile

import pytest
from docx import Document

from helpers import load_jd, load_profile
from resume_builder.ats_check import extract_docx_lines
from resume_builder.matcher import select
from resume_builder.pipeline import full_resume
from resume_builder.render import docx_to_pdf, render_docx
from resume_builder.render.pdf import find_soffice
from resume_builder.writer import base_resume


def _resume():
    p, jd = load_profile(), load_jd()
    return base_resume(p, jd, select(p, jd))


def test_docx_has_no_tables_images_or_header_text(tmp_path):
    path = render_docx(_resume(), tmp_path / "r.docx")
    doc = Document(str(path))
    assert len(doc.tables) == 0
    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    for bad in ("<w:drawing", "<w:pict", "<w:txbxContent"):
        assert bad not in xml
    for s in doc.sections:
        assert not any(p.text.strip() for p in s.header.paragraphs)
        assert not any(p.text.strip() for p in s.footer.paragraphs)


def test_real_word_styles_and_standard_headings(tmp_path):
    path = render_docx(_resume(), tmp_path / "r.docx")
    doc = Document(str(path))
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Title", "Heading 1", "List Bullet", "Normal"} <= styles
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["Summary", "Skills", "Experience", "Projects", "Education", "Certifications"]


def test_header_block_in_body_first_lines(tmp_path):
    r = _resume()
    lines = extract_docx_lines(render_docx(r, tmp_path / "r.docx"))
    assert lines[0] == "Asha Verma"
    assert "asha.verma@example.com" in lines[1] and "linkedin.com/in/ashaverma" in lines[1]


def test_job_entry_line_format(tmp_path):
    lines = extract_docx_lines(render_docx(_resume(), tmp_path / "r.docx"))
    assert "Software Engineer · Acme Corp · Pune · Jun 2021 – Present" in lines
    assert "Junior Developer · Zenith Labs · Ahmedabad · Jul 2019 – May 2021" in lines


def test_early_career_puts_education_before_experience(tmp_path):
    r = _resume()
    r.career_stage = "early_career"
    lines = extract_docx_lines(render_docx(r, tmp_path / "r.docx"))
    assert lines.index("Education") < lines.index("Experience")


def test_fonts_are_calibri(tmp_path):
    path = render_docx(_resume(), tmp_path / "r.docx")
    xml = zipfile.ZipFile(path).read("word/styles.xml").decode()
    normal = re.search(r'<w:style [^>]*w:styleId="Normal".*?</w:style>', xml, re.S).group(0)
    assert "Calibri" in normal


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_pdf_text_extracts_in_order_with_clean_bullets(tmp_path):
    from resume_builder.ats_check import GARBLED, extract_pdf_text

    r = _resume()
    pdf = docx_to_pdf(render_docx(r, tmp_path / "r.docx"), tmp_path / "r.pdf")
    text = extract_pdf_text(pdf)
    assert not GARBLED.search(text), "private-use / ligature glyphs in PDF text"
    order = [text.index(h) for h in ("Summary", "Skills", "Experience", "Projects", "Education", "Certifications")]
    assert order == sorted(order)
    assert "•" in text


def test_full_profile_render(tmp_path):
    p = load_profile()
    lines = extract_docx_lines(render_docx(full_resume(p), tmp_path / "full.docx"))
    assert sum(1 for l in lines if l.startswith("- ")) >= len(p.all_bullets())


@pytest.mark.skipif(find_soffice() is None, reason="LibreOffice not installed")
def test_long_education_line_does_not_split_dates_in_pdf(tmp_path):
    from resume_builder.ats_check import DATE_RANGE, extract_pdf_text

    r = _resume()
    r.education[0].institution = "A Very Long Institute of Engineering and Technology and Research"
    pdf = docx_to_pdf(render_docx(r, tmp_path / "r.docx"), tmp_path / "r.pdf")
    lines = [l for l in extract_pdf_text(pdf).split("\n") if l.strip()]
    assert any(DATE_RANGE.search(l) and "2015" in l for l in lines)
