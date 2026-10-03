"""Step 6a: Resume -> DOCX. Real Word styles (Title, Heading 1, List Bullet, Normal), no tables/images/headers."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ..models import Resume
from ..textutil import date_range

# ---- style constants ("Plain Chronological v1")
FONT = "Calibri"
BODY_PT = 10.0
NAME_PT = 18
HEADING_PT = 12
MARGIN_IN = 0.8
PAPER_IN = {"Letter": (8.5, 11.0), "A4": (8.27, 11.69)}
SEP = " · "


def _set_style_font(style, size: float, bold: bool | None = None) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in list(rfonts.attrib):  # drop theme font overrides so the explicit font wins
        del rfonts.attrib[attr]
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def _remove_border(style) -> None:
    ppr = style.element.get_or_add_pPr()
    for b in ppr.findall(qn("w:pBdr")):
        ppr.remove(b)


def _bottom_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def _round_bullets(doc) -> None:
    """Force a plain round bullet in a standard font (the default template uses Symbol, which extracts as a
    private-use glyph in some PDF text extractors)."""
    try:
        numbering = doc.part.numbering_part.element
    except Exception:
        return
    for lvl in numbering.iter(qn("w:lvl")):
        fmt = lvl.find(qn("w:numFmt"))
        if fmt is None or fmt.get(qn("w:val")) != "bullet":
            continue
        text = lvl.find(qn("w:lvlText"))
        if text is not None:
            text.set(qn("w:val"), "•")
        rpr = lvl.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            lvl.append(rpr)
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in list(rfonts.attrib):
            del rfonts.attrib[attr]
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), FONT)


def _setup(doc: Document, resume: Resume) -> None:
    sec = doc.sections[0]
    w, h = PAPER_IN[resume.paper_size]
    sec.page_width, sec.page_height = Inches(w), Inches(h)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(MARGIN_IN))

    normal = doc.styles["Normal"]
    _set_style_font(normal, BODY_PT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    title = doc.styles["Title"]
    _set_style_font(title, NAME_PT, bold=True)
    _remove_border(title)
    title.paragraph_format.space_after = Pt(2)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    h1 = doc.styles["Heading 1"]
    _set_style_font(h1, HEADING_PT, bold=True)
    h1.paragraph_format.space_before = Pt(7)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.keep_with_next = True

    lb = doc.styles["List Bullet"]
    _set_style_font(lb, BODY_PT)
    lb.paragraph_format.space_after = Pt(1)
    _round_bullets(doc)

    doc.core_properties.author = resume.basics.name
    doc.core_properties.title = f"{resume.basics.name} - Resume"
    doc.core_properties.comments = ""


def _entry_line(doc, bold: str, rest: list[str]) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.add_run(bold).bold = True
    tail = [r for r in rest if r]
    if tail:
        p.add_run(SEP + SEP.join(tail))


def _bullets(doc, texts: list[str]) -> None:
    for t in texts:
        # one consistent style: no trailing period (keeps 'Node.js' safe because it ends in 's')
        doc.add_paragraph(t.rstrip().rstrip(".") if not t.rstrip().endswith("..") else t, style="List Bullet")


def section_order(resume: Resume) -> list[str]:
    if resume.career_stage == "early_career":
        return ["Summary", "Skills", "Education", "Experience", "Projects", "Certifications"]
    return ["Summary", "Skills", "Experience", "Projects", "Education", "Certifications"]


def render_docx(resume: Resume, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    _setup(doc, resume)
    fmt = resume.date_format

    # header block lives in the document body (not the Word header) so every parser sees it
    b = resume.basics
    doc.add_paragraph(b.name, style="Title")
    contact = [x for x in [b.location, b.phone, b.email, *[l.url for l in b.links]] if x]
    doc.add_paragraph(" | ".join(contact))

    def heading(name: str) -> None:
        _bottom_border(doc.add_paragraph(name, style="Heading 1"))

    for name in section_order(resume):
        if name == "Summary" and resume.summary:
            heading(name)
            doc.add_paragraph(resume.summary)
        elif name == "Skills" and resume.skills:
            heading(name)
            for g in resume.skills:
                p = doc.add_paragraph()
                p.add_run(f"{g.category}: ").bold = True
                p.add_run(", ".join(g.skills))
        elif name == "Experience" and resume.experience:
            heading(name)
            for e in resume.experience:
                _entry_line(doc, e.title, [e.company, e.location, date_range(e.start, e.end, fmt)])
                _bullets(doc, [x.text for x in e.bullets])
        elif name == "Projects" and resume.projects:
            heading(name)
            for pr in resume.projects:
                _entry_line(doc, pr.name, [pr.url])
                _bullets(doc, [x.text for x in pr.bullets])
        elif name == "Education" and resume.education:
            heading(name)
            for ed in resume.education:
                dates = date_range(ed.start, ed.end, fmt) if ed.start or ed.end else ""
                one_line = SEP.join(x for x in [ed.degree, ed.institution, ed.location, dates] if x)
                if len(one_line) > 95:  # a wrapped line can split the date range, so break it deliberately
                    _entry_line(doc, ed.degree, [ed.institution])
                    doc.add_paragraph(SEP.join(x for x in [ed.location, dates] if x)).paragraph_format.keep_with_next = True
                else:
                    _entry_line(doc, ed.degree, [ed.institution, ed.location, dates])
                # Render education details as plain paragraphs (not list bullets) so
                # external ATS checkers don't flag "CGPA: x.xx/10" as a bullet without a verb
                for detail in ed.details:
                    p = doc.add_paragraph(detail)
                    p.paragraph_format.left_indent = Inches(0.15)
                    p.paragraph_format.space_after = Pt(1)
        elif name == "Certifications" and resume.certifications:
            heading(name)
            for c in resume.certifications:
                doc.add_paragraph(SEP.join(x for x in [c.name, c.issuer, c.date] if x))
    doc.save(path)
    return path
