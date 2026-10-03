"""Step 7: re-read the OUTPUT FILE the way a parser would and score parse quality (a proxy, not a promise)."""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path

from .models import JD, Resume
from .textutil import STANDARD_HEADINGS, SkillIndex, date_range, term_in_text

WEIGHTS = {"Parse fidelity": 40, "Keyword coverage": 35, "Format compliance": 15, "Readability": 10}
_MON = r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)"
DATE_RANGE = re.compile(rf"(?:{_MON}\.? )?\d{{4}}\s*[–\-]\s*(?:(?:{_MON}\.? )?\d{{4}}|Present)", re.I)
EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"\+?\d[\d\s().-]{7,}\d")
ALLOWED_FONTS = {"calibri", "arial"}
IRREGULAR_VERBS = {
    "led", "built", "ran", "drove", "wrote", "cut", "grew", "won", "made", "set", "took", "sped", "began",
    "taught", "oversaw", "spun", "sold", "brought", "kept", "held", "chose", "saw", "put", "found",
}
GARBLED = re.compile("[\ue000-\uf8ff\ufb00-\ufb06\ufffd]")


@dataclass
class Check:
    name: str
    earned: float
    weight: float
    detail: str = ""


@dataclass
class ATSReport:
    components: dict[str, list[Check]] = field(default_factory=dict)
    not_applicable: dict[str, str] = field(default_factory=dict)
    matched_must: list[str] = field(default_factory=list)
    missing_must: list[str] = field(default_factory=list)
    matched_nice: list[str] = field(default_factory=list)
    missing_nice: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    text: str = ""
    pages: int | None = None

    def score(self, component: str) -> float:
        checks = self.components.get(component, [])
        total = sum(c.weight for c in checks)
        return WEIGHTS[component] * sum(c.earned * c.weight for c in checks) / total if total else 0.0

    def max_points(self, component: str) -> float:
        return 0 if component in self.not_applicable else WEIGHTS[component]

    @property
    def overall(self) -> float:
        applicable = [c for c in WEIGHTS if c not in self.not_applicable and self.components.get(c)]
        denom = sum(WEIGHTS[c] for c in applicable)
        return round(100 * sum(self.score(c) for c in applicable) / denom, 1) if denom else 0.0

    def summary_line(self) -> str:
        parts = []
        for c, short in (("Parse fidelity", "Parse"), ("Keyword coverage", "Keywords"), ("Format compliance", "Format"), ("Readability", "Readability")):
            if c in self.not_applicable:
                parts.append(f"{short} n/a")
            else:
                parts.append(f"{short} {self.score(c):.0f}/{WEIGHTS[c]}")
        return f"Overall: {self.overall:.0f}/100  ({' · '.join(parts)})"


# ------------------------------------------------------------------ extraction
def extract_docx_lines(path: str | Path) -> list[str]:
    """Body text in document order. Bullets are marked '- ' (Word stores them as numbering, not text)."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph

            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = para.style.name if para.style is not None else ""
            lines.append(("- " + text) if style.startswith("List") else text)
        elif child.tag == qn("w:tbl"):
            for t in child.iter(qn("w:t")):
                if t.text and t.text.strip():
                    lines.append(t.text.strip())
    return lines


def extract_pdf_text(path: str | Path) -> str:
    import pdfplumber

    with pdfplumber.open(str(path)) as pdf:
        return "\n".join((page.extract_text() or "") for page in pdf.pages)


# ------------------------------------------------------------------ helpers
def _heading_positions(lines: list[str]) -> dict[str, int]:
    pos: dict[str, int] = {}
    for i, l in enumerate(lines):
        for h in STANDARD_HEADINGS:
            if l.strip().lower() == h.lower() and h not in pos:
                pos[h] = i
    return pos


def _expected_headings(resume: Resume) -> list[str]:
    from .render.docx_renderer import section_order

    have = {
        "Summary": bool(resume.summary), "Skills": bool(resume.skills), "Experience": bool(resume.experience),
        "Projects": bool(resume.projects), "Education": bool(resume.education),
        "Certifications": bool(resume.certifications),
    }
    return [h for h in section_order(resume) if have[h]]


def _bullets(lines: list[str]) -> list[str]:
    return [re.sub(r"^[-•\u2022]\s*", "", l) for l in lines if re.match(r"^[-•\u2022]\s+\S", l)]


def _docx_xml(path: Path) -> tuple[str, str]:
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf8"), z.read("word/styles.xml").decode("utf8")


# ------------------------------------------------------------------ the check
def check_file(
    path: str | Path,
    *,
    resume: Resume | None = None,
    jd: JD | None = None,
    idx: SkillIndex | None = None,
    max_pages: int = 1,
    pdf_path: str | Path | None = None,
) -> ATSReport:
    """`path` is the DOCX or PDF to inspect. Give `resume` for exact checks; without it, generic checks run."""
    path = Path(path)
    rep = ATSReport()
    is_docx = path.suffix.lower() == ".docx"
    if is_docx:
        lines = extract_docx_lines(path)
        pdf = Path(pdf_path) if pdf_path and Path(pdf_path).exists() else None
    else:
        pdf = path
        lines = [l for l in extract_pdf_text(path).split("\n") if l.strip()]
    pdf_text = extract_pdf_text(pdf) if pdf else ""
    pdf_lines = [l.strip() for l in pdf_text.split("\n") if l.strip()]
    if pdf:
        from .render.pdf import page_count

        rep.pages = page_count(pdf)
    rep.text = "\n".join(lines)
    full = rep.text

    _parse_fidelity(rep, lines, pdf_lines if pdf else None, resume)
    _keywords(rep, full, jd, idx)
    _format(rep, path if is_docx else None, pdf, max_pages)
    _readability(rep, lines, jd)
    return rep


def _parse_fidelity(rep: ATSReport, lines: list[str], pdf_lines: list[str] | None, resume: Resume | None) -> None:
    checks: list[Check] = []
    top = lines[:8]
    top_text = "\n".join(top)

    # 1. contact block at the top
    if resume:
        name_ok = any(resume.basics.name.lower() in l.lower() for l in lines[:3])
        email_ok = resume.basics.email.lower() in top_text.lower()
        phone_ok = (not resume.basics.phone) or re.sub(r"\D", "", resume.basics.phone) in re.sub(r"\D", "", top_text)
        checks.append(Check("name in first lines", float(name_ok), 5, "" if name_ok else "name not found in first 3 lines"))
        checks.append(Check("email in contact block", float(email_ok), 3, "" if email_ok else "email not in first 8 lines"))
        checks.append(Check("phone in contact block", float(phone_ok), 2, "" if phone_ok else "phone not in first 8 lines"))
    else:
        first = lines[0] if lines else ""
        name_ok = bool(first) and "@" not in first and len(first.split()) <= 5 and not first.startswith("-")
        checks.append(Check("name-like first line", float(name_ok), 5, "" if name_ok else f"first line is '{first[:40]}'"))
        checks.append(Check("email in first lines", float(bool(EMAIL.search(top_text))), 3, "no email in first 8 lines"))
        checks.append(Check("phone in first lines", float(bool(PHONE.search(top_text))), 2, "no phone in first 8 lines"))

    # 2. standard headings, in order
    pos = _heading_positions(lines)
    if resume:
        expected = _expected_headings(resume)
        found = [h for h in expected if h in pos]
        ordered = [h for h in expected if h in pos] == sorted(found, key=lambda h: pos[h])
        checks.append(Check("standard headings found", len(found) / len(expected) if expected else 1, 7,
                            "missing: " + ", ".join(h for h in expected if h not in pos) if len(found) != len(expected) else ""))
        checks.append(Check("headings in expected order", float(ordered and bool(found)), 3, "" if ordered else "headings out of order"))
    else:
        n = len(pos)
        checks.append(Check("standard headings found", min(1.0, n / 4), 7, f"found {n}: {', '.join(pos)}" if n < 4 else ""))
        canon = [h for h in ("Summary", "Skills", "Education", "Experience", "Projects", "Certifications") if h in pos]
        seq = sorted(pos, key=lambda h: pos[h])
        ok = seq == canon or seq == [h for h in ("Summary", "Skills", "Experience", "Projects", "Education", "Certifications") if h in pos]
        checks.append(Check("headings in a conventional order", float(ok), 3, "" if ok else "unusual section order"))

    # 3. job entries: title, company, dates on one line, in order
    if resume and resume.experience:
        from .render.docx_renderer import SEP

        hit, idxs, bad = 0, [], []
        for e in resume.experience:
            dr = date_range(e.start, e.end, resume.date_format)
            line_no = next((i for i, l in enumerate(lines) if e.title in l and e.company in l and dr in l), None)
            if line_no is None:
                bad.append(e.company)
            else:
                hit += 1
                idxs.append(line_no)
        in_order = idxs == sorted(idxs)
        checks.append(Check("job entries parse (title/company/dates)", hit / len(resume.experience), 8, "not parsed: " + ", ".join(bad) if bad else ""))
        checks.append(Check("job entries in reverse-chronological order", float(in_order), 2, "" if in_order else "entries out of order"))
    else:
        n = sum(1 for l in lines if DATE_RANGE.search(l) and len(l) > 20)
        checks.append(Check("date-range entries detected", float(n > 0), 10, "no 'Mon YYYY – Mon YYYY' lines found" if n == 0 else ""))

    # 4. garbled characters
    all_text = "\n".join(lines) + "\n" + "\n".join(pdf_lines or [])
    bad_chars = sorted({f"U+{ord(c):04X}" for c in GARBLED.findall(all_text)})
    checks.append(Check("no garbled/ligature/private-use characters", float(not bad_chars), 4, ", ".join(bad_chars)))

    # 5. reading order, DOCX vs PDF
    if pdf_lines is None:
        checks.append(Check("reading order (DOCX vs PDF)", 0.5, 6, "no PDF available: order not cross-checked"))
        rep.warnings.append("PDF not available: reading-order cross-check skipped")
    else:
        d = sorted(_heading_positions(lines), key=lambda h: _heading_positions(lines)[h])
        p = sorted(_heading_positions(pdf_lines), key=lambda h: _heading_positions(pdf_lines)[h])
        heads_ok = d == p and bool(d)
        d_dates = [m.group(0) for l in lines for m in [DATE_RANGE.search(l)] if m]
        p_dates = [m.group(0) for l in pdf_lines for m in [DATE_RANGE.search(l)] if m]
        dates_ok = d_dates == p_dates
        checks.append(Check("heading order matches in PDF", float(heads_ok), 3, "" if heads_ok else f"DOCX {d} vs PDF {p}"))
        checks.append(Check("entry order matches in PDF", float(dates_ok), 3, "" if dates_ok else "date lines differ between DOCX and PDF"))
    rep.components["Parse fidelity"] = checks


def _keywords(rep: ATSReport, text: str, jd: JD | None, idx: SkillIndex | None) -> None:
    if jd is None or not (jd.must_have or jd.nice_to_have):
        rep.not_applicable["Keyword coverage"] = "no job description supplied"
        return
    has = (lambda t: idx.text_has(t, text)) if idx else (lambda t: term_in_text(t, text))
    rep.matched_must = [t for t in jd.must_have if has(t)]
    rep.missing_must = [t for t in jd.must_have if not has(t)]
    rep.matched_nice = [t for t in jd.nice_to_have if has(t)]
    rep.missing_nice = [t for t in jd.nice_to_have if not has(t)]
    checks = []
    if jd.must_have:
        cov = len(rep.matched_must) / len(jd.must_have)
        checks.append(Check("must-have coverage", cov, 70, f"missing: {', '.join(rep.missing_must)}" if rep.missing_must else ""))
    if jd.nice_to_have:
        cov = len(rep.matched_nice) / len(jd.nice_to_have)
        checks.append(Check("nice-to-have coverage", cov, 30, f"missing: {', '.join(rep.missing_nice)}" if rep.missing_nice else ""))
    rep.components["Keyword coverage"] = checks


def _format(rep: ATSReport, docx: Path | None, pdf: Path | None, max_pages: int) -> None:
    checks: list[Check] = []
    if docx is not None:
        from docx import Document

        doc_xml, styles_xml = _docx_xml(docx)
        doc = Document(str(docx))
        checks.append(Check("no tables", float(len(doc.tables) == 0), 3, f"{len(doc.tables)} table(s)" if doc.tables else ""))
        shapes = [t for t in ("<w:drawing", "<w:pict", "<w:txbxContent", "<mc:AlternateContent", "<v:shape") if t in doc_xml]
        checks.append(Check("no images/text boxes/shapes", float(not shapes), 3, ", ".join(shapes)))
        hf = []
        for s in doc.sections:
            for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
                if part is not None and not part.is_linked_to_previous and any(p.text.strip() for p in part.paragraphs):
                    hf.append("header/footer text")
        checks.append(Check("no header/footer content", float(not hf), 3, "header/footer contains text" if hf else ""))
        fonts = _fonts_used(doc, doc_xml)
        bad = sorted(f for f in fonts if f.lower() not in ALLOWED_FONTS)
        checks.append(Check("standard fonts (Calibri/Arial)", float(not bad), 2, ", ".join(bad)))
        cols = re.findall(r'<w:cols[^>]*w:num="(\d+)"', doc_xml)
        multi = any(int(c) > 1 for c in cols)
        checks.append(Check("single column", float(not multi), 2, "multi-column section" if multi else ""))
    else:
        rep.warnings.append("Input is a PDF: DOCX structure checks (tables, headers, fonts) were skipped")

    if rep.pages is not None:
        ok = rep.pages <= max_pages
        checks.append(Check(f"page count <= {max_pages}", float(ok), 2, f"{rep.pages} pages" if not ok else ""))
    else:
        checks.append(Check("page count", 0.5, 2, "no PDF to measure"))
    if pdf is not None:
        has_text = len(extract_pdf_text(pdf).split()) > 30
        checks.append(Check("PDF has selectable text", float(has_text), 2, "" if has_text else "PDF looks image-only"))
    rep.components["Format compliance"] = checks


def _fonts_used(doc, doc_xml: str) -> set[str]:
    fonts = set(re.findall(r'<w:rFonts[^>]*w:ascii="([^"]+)"', doc_xml))
    seen = set()
    for p in doc.paragraphs:
        st = p.style
        while st is not None and st.style_id not in seen:
            seen.add(st.style_id)
            if st.font is not None and st.font.name:
                fonts.add(st.font.name)
                break
            st = st.base_style
    return fonts


def _similar(a: str, b: str) -> float:
    try:
        from rapidfuzz import fuzz

        return fuzz.ratio(a.lower(), b.lower()) / 100
    except ImportError:  # pragma: no cover
        return SequenceMatcher(None, a.lower(), b.lower()).ratio()


def _readability(rep: ATSReport, lines: list[str], jd: JD | None) -> None:
    bullets = _bullets(lines)
    if not bullets:
        rep.components["Readability"] = [Check("bullets present", 0.0, 10, "no bullets found")]
        return
    checks: list[Check] = []
    long_ = [b for b in bullets if len(b) > 210]
    checks.append(Check("bullets <= 2 lines", 1 - len(long_) / len(bullets), 3, f"{len(long_)} bullet(s) exceed 2 lines" if long_ else ""))
    if long_:
        rep.warnings.append(f"{len(long_)} bullet(s) exceed 2 lines")

    def verb_first(b: str) -> bool:
        w = re.sub(r"[^A-Za-z]", "", b.split()[0]).lower() if b.split() else ""
        return w.endswith("ed") or w in IRREGULAR_VERBS

    weak = [b for b in bullets if not verb_first(b)]
    checks.append(Check("bullets start with a verb", 1 - len(weak) / len(bullets), 3, f"{len(weak)} bullet(s) do not start with a verb" if weak else ""))
    if weak:
        rep.warnings.append(f"{len(weak)} bullet(s) may not start with an action verb")

    dups = sum(1 for i, a in enumerate(bullets) for b in bullets[i + 1:] if _similar(a, b) > 0.9)
    checks.append(Check("no duplicate bullets", float(dups == 0), 2, f"{dups} near-duplicate pair(s)" if dups else ""))

    with_num = [b for b in bullets if re.search(r"\d", b)]
    ratio = len(with_num) / len(bullets)
    checks.append(Check("bullets with a number/metric (aim for >= 50%)", min(1.0, ratio / 0.5), 3,
                        f"{len(with_num)} of {len(bullets)} bullets contain a number" if ratio < 0.5 else ""))
    if ratio < 0.5:
        rep.warnings.append(f"only {len(with_num)} of {len(bullets)} bullets contain a number: add real metrics to profile.json")

    verbs: dict[str, int] = {}
    for b in bullets:
        w = re.sub(r"[^A-Za-z]", "", b.split()[0]).lower() if b.split() else ""
        verbs[w] = verbs.get(w, 0) + 1
    repeated = {v: n for v, n in verbs.items() if n > 2}
    excess = sum(n - 2 for n in repeated.values())
    checks.append(Check("varied opening verbs (max 2 per verb)", max(0.0, 1 - excess / len(bullets)), 2,
                        ", ".join(f"{v} x{n}" for v, n in repeated.items())))
    if repeated:
        rep.warnings.append("repeated opening verbs: " + ", ".join(f"{v} x{n}" for v, n in repeated.items()))

    ends = {b.rstrip().endswith(".") for b in bullets}
    checks.append(Check("consistent bullet punctuation", float(len(ends) == 1), 1, "" if len(ends) == 1 else "some bullets end with a period, others do not"))

    stuffed: list[str] = []
    if jd:
        for t in [*jd.must_have, *jd.nice_to_have, *jd.keywords]:
            if t and len(term_regex_findall(t, "\n".join(lines))) > 5:
                stuffed.append(t)
    checks.append(Check("no keyword stuffing", float(not stuffed), 2, "repeated >5x: " + ", ".join(stuffed) if stuffed else ""))
    if stuffed:
        rep.warnings.append("possible keyword stuffing: " + ", ".join(stuffed))
    rep.components["Readability"] = checks


def term_regex_findall(term: str, text: str) -> list[str]:
    from .textutil import term_regex

    return term_regex(term).findall(text)
